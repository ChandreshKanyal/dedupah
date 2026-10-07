"""Taking results out: downloads (Excel, CSV, Parquet), writes into a database (new table, add to one, replace one), the destinations they go to, and
the safeguards around personal data leaving: permission, reason, rate limit, audit, formula safety, all-or-nothing database writes."""
import csv
import io

import pyarrow as pa
import pyarrow.parquet as pq
import pytest
from api_world import PROJECT, app_with_run, hdr, pipeline_output
from openpyxl import load_workbook
from sqlalchemy import create_engine, text

from dedupe.results import ExportError, describe, write_database, write_file
from dedupe.results.export import _csv_cell

ROOT_USER = ("admin", "u-root")


@pytest.fixture(scope="module")
def output(tmp_path_factory):
    return pipeline_output(tmp_path_factory.mktemp("export_out"), seed=81, households=40)


def build(tmp_path, output, monkeypatch, **limits):
    monkeypatch.setenv("DEDUPE_DST_WAREHOUSE", f"sqlite:///{tmp_path / 'warehouse.db'}")
    c, s, sv, run = app_with_run(tmp_path, output, ip_rate_limit_per_min=10**6, rate_limit_per_min=10**6, **limits)
    return c, s, sv, run, tmp_path


@pytest.fixture
def world(tmp_path, output, monkeypatch):
    return build(tmp_path, output, monkeypatch, export_limit_per_min=10**6)           # the many calls of these tests are not what the limit is for


@pytest.fixture
def limited(tmp_path, output, monkeypatch):
    return build(tmp_path, output, monkeypatch, export_limit_per_min=6)


def make_destination_live(c, s, name="warehouse", env="DEDUPE_DST_WAREHOUSE", schema=None):
    body = {"destination": {"destination_id": name, "url_env": env, "schema_name": schema}, "reason": "the reporting warehouse"}
    r = c.post(f"/projects/{PROJECT}/destinations", json=body, headers=hdr(s, "platform_admin", "u-ops"))
    assert r.status_code == 200, r.text
    base = f"/projects/{PROJECT}/destinations/{name}/versions/{r.json()['version']}"
    assert c.post(base + "/approve", json={}, headers=hdr(s, "approver", "u-lee")).status_code == 200
    assert c.post(base + "/activate", json={}, headers=hdr(s, "approver", "u-lee")).status_code == 200


def rows_in(db, table):
    with create_engine(f"sqlite:///{db}").connect() as conn:
        return conn.execute(text(f"SELECT count(*) FROM {table}")).scalar()


def test_a_run_offers_its_datasets_with_person_ids_first_and_lists_as_text(world):
    c, s, sv, run, tmp = world
    info = c.get(run + "/datasets", headers=hdr(s, *ROOT_USER)).json()
    by = {d["name"]: d for d in info["datasets"]}
    assert set(by) == {"golden_records", "person_records", "households"} and info["formats"] == {"xlsx": "Excel", "csv": "CSV", "parquet": "Parquet"}
    assert by["golden_records"]["columns"][0] == "person_id" and "cluster_id" not in by["golden_records"]["columns"] and by["golden_records"]["rows"] > 10
    assert by["person_records"]["columns"] == ["person_id", "source_id", "source_record_id"]                      # the internal row number is not offered
    assert by["households"]["columns"][0] == "person_id" and "household_id" in by["households"]["columns"]
    assert by["golden_records"]["rows"] == describe(sv.workspace.run_dir(PROJECT, run.rsplit("/", 1)[1]))[0]["rows"]


@pytest.mark.parametrize("fmt", ["csv", "xlsx", "parquet"])
def test_each_format_downloads_the_same_rows(world, fmt):
    c, s, sv, run, tmp = world
    rows = {d["name"]: d["rows"] for d in c.get(run + "/datasets", headers=hdr(s, *ROOT_USER)).json()["datasets"]}
    r = c.post(run + "/export", json={"dataset": "golden_records", "format": fmt, "reason": "monthly report to the board"}, headers=hdr(s, *ROOT_USER))
    assert r.status_code == 200 and f"{PROJECT}_" in r.headers["content-disposition"] and r.headers["content-disposition"].endswith(f'.{fmt}"')
    assert r.headers["cache-control"] == "no-store"
    if fmt == "csv":
        data = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig"))))
        header, body = data[0], data[1:]
    elif fmt == "xlsx":
        ws = load_workbook(io.BytesIO(r.content), read_only=True).active
        data = [list(x) for x in ws.iter_rows(values_only=True)]
        header, body = data[0], data[1:]
    else:
        table = pq.read_table(io.BytesIO(r.content))
        header, body = table.column_names, table.to_pylist()
    assert header[0] == "person_id" and len(body) == rows["golden_records"]
    assert not list((sv.settings.work / "exports").glob("*"))                                                    # the temporary file is deleted after sending


def test_exports_need_the_permission_a_reason_and_a_finished_run_you_can_open(world):
    c, s, sv, run, tmp = world
    good = {"dataset": "households", "format": "csv", "reason": "needed for the mailing list"}
    assert c.post(run + "/export", json=good, headers=hdr(s, "viewer")).status_code == 403                  # reading is not taking out
    assert c.post(run + "/export", json=good, headers=hdr(s, "steward")).status_code == 403
    assert c.post(run + "/export", json={**good, "reason": "short"}, headers=hdr(s, *ROOT_USER)).status_code == 422
    assert c.post(run + "/export", json={**good, "dataset": "everything"}, headers=hdr(s, *ROOT_USER)).status_code == 422
    assert c.post(run + "/export", json={**good, "format": "pdf"}, headers=hdr(s, *ROOT_USER)).status_code == 422
    assert c.post(f"/projects/{PROJECT}/runs/R000000000000/export", json=good, headers=hdr(s, *ROOT_USER)).status_code == 404
    c.post(f"/projects/{PROJECT}/access", json={"subject": "u-root", "role": "admin"}, headers=hdr(s, *ROOT_USER))     # the project becomes members-only
    assert c.post(run + "/export", json=good, headers=hdr(s, "analyst", "outsider")).status_code == 403
    assert c.get(run + "/datasets", headers=hdr(s, "analyst", "outsider")).status_code == 403


def test_every_export_is_audited_with_counts_never_values_and_is_rate_limited(limited):
    c, s, sv, run, tmp = limited
    admin = hdr(s, *ROOT_USER)
    r = c.post(run + "/export", json={"dataset": "golden_records", "format": "csv", "reason": "monthly report to the board"}, headers=admin)
    assert r.status_code == 200
    ev = [e for e in sv.seclog.events(limit=500, event_name="export") if e["detail"].get("action") == "export_file"][-1]
    assert ev["actor"] == "u-root" and ev["project_id"] == PROJECT and ev["detail"]["dataset"] == "golden_records" and ev["detail"]["format"] == "csv"
    assert ev["detail"]["rows"] > 10 and ev["detail"]["reason"] == "monthly report to the board"
    first_name = r.content.decode("utf-8-sig").splitlines()[1].split(",")[3]
    assert first_name and first_name not in str(ev)                                                                # no value from the file is in the log
    codes = [c.post(run + "/export", json={"dataset": "households", "format": "csv", "reason": "load test of the limit"}, headers=admin).status_code for _ in range(8)]
    assert 429 in codes and codes[0] == 200


def test_values_that_a_spreadsheet_would_run_as_a_formula_are_made_safe(tmp_path):
    run_dir = tmp_path / "run"
    (run_dir / "golden").mkdir(parents=True)
    pq.write_table(pa.table({"cluster_id": ["P000000001", "P000000002"], "record_count": [2, 1], "sources": [["a", "b"], ["a"]], "full_name": ["=HYPERLINK(\"http://x\",\"y\")", "Asha"],
                             "mobile": ["+919876543210", "9876543210"], "conflict_fields": [[], ["pan"]]}), run_dir / "golden" / "golden_records.parquet")
    assert _csv_cell("=1+1", True) == "'=1+1" and _csv_cell("=1+1", False) == "=1+1" and _csv_cell("Asha", True) == "Asha" and _csv_cell(None, True) == ""
    write_file(run_dir, "golden_records", "csv", tmp_path / "o.csv")
    text_out = (tmp_path / "o.csv").read_text(encoding="utf-8-sig")
    assert "'=HYPERLINK" in text_out and "'+919876543210" in text_out and "a; b" in text_out                       # lists become text; risky values get an apostrophe
    write_file(run_dir, "golden_records", "csv", tmp_path / "raw.csv", safe_for_spreadsheets=False)
    assert "\n+919876543210" not in (tmp_path / "raw.csv").read_text(encoding="utf-8-sig") and ",+919876543210" in (tmp_path / "raw.csv").read_text(encoding="utf-8-sig")
    write_file(run_dir, "golden_records", "xlsx", tmp_path / "o.xlsx")
    cell = next(iter(load_workbook(tmp_path / "o.xlsx").active.iter_rows(min_row=2, max_row=2)))[3]
    assert cell.data_type == "s" and cell.value.startswith("=HYPERLINK")                                            # stored as text, never evaluated


def test_a_database_destination_is_saved_checked_and_approved_like_a_connection(world, monkeypatch):
    c, s, sv, run, tmp = world
    ops = hdr(s, "platform_admin", "u-ops")
    base = f"/projects/{PROJECT}/destinations"
    unset = c.post(base, json={"destination": {"destination_id": "nowhere", "url_env": "DEDUPE_DST_NOT_SET"}, "reason": "the reporting warehouse"}, headers=ops)
    assert unset.status_code == 422 and "not set" in unset.json()["error"]["message"]
    assert c.post(base, json={"destination": {"destination_id": "w", "url_env": "DATABASE_URL"}, "reason": "the reporting warehouse"}, headers=ops).status_code == 422
    assert c.post(base, json={"destination": {"destination_id": "w", "url_env": "DEDUPE_DST_WAREHOUSE", "password": "x"}, "reason": "the reporting warehouse"}, headers=ops).status_code == 422
    monkeypatch.setenv("DEDUPE_DST_BROKEN", "sqlite:////no/such/dir/hunter2.db")
    broken = c.post(base, json={"destination": {"destination_id": "b", "url_env": "DEDUPE_DST_BROKEN"}, "reason": "the reporting warehouse"}, headers=ops)
    assert broken.status_code == 422 and "hunter2" not in broken.text and "could not be reached" in broken.json()["error"]["message"]
    ok = c.post(base, json={"destination": {"destination_id": "warehouse", "url_env": "DEDUPE_DST_WAREHOUSE"}, "reason": "the reporting warehouse"}, headers=ops)
    assert ok.status_code == 200 and ok.json()["status"] == "proposed"
    ver = ok.json()["version"]
    export = {"destination": "warehouse", "dataset": "households", "table": "hh", "mode": "create", "reason": "load the household table"}
    assert c.post(run + "/export/database", json=export, headers=hdr(s, *ROOT_USER)).status_code == 409           # not live yet
    assert c.post(f"{base}/warehouse/versions/{ver}/approve", json={}, headers=ops).status_code in (403, 409)      # not by the proposer
    assert c.post(f"{base}/warehouse/versions/{ver}/approve", json={}, headers=hdr(s, "approver", "u-lee")).status_code == 200
    assert c.post(f"{base}/warehouse/versions/{ver}/activate", json={}, headers=hdr(s, "approver", "u-lee")).status_code == 200
    listed = c.get(base, headers=hdr(s, *ROOT_USER)).json()
    assert listed == [{**listed[0], "id": "warehouse", "url_env": "DEDUPE_DST_WAREHOUSE", "live_version": 1}]
    assert c.post(base + "/warehouse/test", headers=ops).json() == {"ok": True, "error": None}
    assert c.get(base, headers=hdr(s, "viewer")).status_code == 403                                                 # only people who may export see destinations


def test_database_writes_create_append_and_replace_a_table(world):
    c, s, sv, run, tmp = world
    make_destination_live(c, s)
    admin = hdr(s, *ROOT_USER)
    db = tmp / "warehouse.db"
    n = next(d["rows"] for d in c.get(run + "/datasets", headers=admin).json()["datasets"] if d["name"] == "households")
    body = {"destination": "warehouse", "dataset": "households", "table": "dedupe_households", "mode": "create", "reason": "load the household table"}
    r = c.post(run + "/export/database", json=body, headers=admin)
    assert r.status_code == 200 and r.json()["rows"] == n and rows_in(db, "dedupe_households") == n
    again = c.post(run + "/export/database", json=body, headers=admin)
    assert again.status_code == 422 and "already exists" in again.json()["error"]["message"]                         # create never overwrites
    assert c.post(run + "/export/database", json={**body, "mode": "append"}, headers=admin).status_code == 200
    assert rows_in(db, "dedupe_households") == 2 * n                                                                  # appended below the first load
    unconfirmed = c.post(run + "/export/database", json={**body, "mode": "replace"}, headers=admin)
    assert unconfirmed.status_code == 422 and "confirm" in unconfirmed.json()["error"]["message"] and rows_in(db, "dedupe_households") == 2 * n
    assert c.post(run + "/export/database", json={**body, "mode": "replace", "confirm_replace": True}, headers=admin).status_code == 200
    assert rows_in(db, "dedupe_households") == n                                                                       # replaced: back to one load
    assert c.post(run + "/export/database", json={**body, "table": "missing_table", "mode": "append"}, headers=admin).status_code == 422
    for bad in ("x; drop table dedupe_households", "1abc", "has space", "a" * 70):
        assert c.post(run + "/export/database", json={**body, "table": bad}, headers=admin).status_code == 422
    assert c.post(run + "/export/database", json={**body, "mode": "replace", "confirm_replace": True}, headers=hdr(s, "viewer")).status_code == 403
    ev = [e for e in sv.seclog.events(limit=500, event_name="export") if e["detail"].get("action") == "export_database"]
    assert len(ev) == 3 and ev[0]["detail"]["table"] == "dedupe_households" and ev[0]["detail"]["destination"] == "warehouse" and ev[0]["detail"]["rows"] == n


def test_a_failed_database_write_leaves_the_table_exactly_as_it_was(world):
    c, s, sv, run, tmp = world
    make_destination_live(c, s)
    admin = hdr(s, *ROOT_USER)
    db = tmp / "warehouse.db"
    eng = create_engine(f"sqlite:///{db}")
    with eng.begin() as conn:
        conn.execute(text("CREATE TABLE hh_needs_more (person_id TEXT, household_id TEXT, household_size INTEGER, region TEXT NOT NULL)"))
        conn.execute(text("INSERT INTO hh_needs_more VALUES ('keep-1', 'x', 1, 'north'), ('keep-2', 'y', 2, 'south')"))
        conn.execute(text("CREATE TABLE hh_narrow (person_id TEXT)"))
    body = {"destination": "warehouse", "dataset": "households", "table": "hh_needs_more", "mode": "replace", "confirm_replace": True, "reason": "load the household table"}
    r = c.post(run + "/export/database", json=body, headers=admin)
    assert r.status_code == 422 and "nothing was written" in r.json()["error"]["message"] and "region" not in r.text                 # the driver's text is not repeated
    with eng.connect() as conn:
        assert [row[0] for row in conn.execute(text("SELECT person_id FROM hh_needs_more ORDER BY person_id"))] == ["keep-1", "keep-2"]      # the delete was rolled back
    narrow = c.post(run + "/export/database", json={**body, "table": "hh_narrow", "mode": "append"}, headers=admin)
    assert narrow.status_code == 422 and "no column for" in narrow.json()["error"]["message"] and rows_in(db, "hh_narrow") == 0


def test_the_engine_refuses_unsafe_requests_on_its_own(tmp_path):
    with pytest.raises(ExportError):
        write_database("sqlite://", tmp_path, "households", "t", "create")                       # the run has no such results
    with pytest.raises(ExportError):
        write_database("sqlite://", tmp_path, "households", "t; drop", "create")
    with pytest.raises(ExportError):
        write_database("sqlite://", tmp_path, "households", "t", "truncate")
    with pytest.raises(ExportError):
        write_file(tmp_path, "everything", "csv", tmp_path / "x.csv")
