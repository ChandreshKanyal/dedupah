"""Files sent from the browser (CSV, TSV, Excel).

An upload is untrusted: the name is reduced to a safe label, the type is decided by the content (not the name), Excel files are opened only after
checking they are not compression bombs, the header must be usable (non-empty, unique, a sensible number of columns), and rows are capped. The
file is converted once into a Parquet table of text columns, stored under a random id (encrypted when a data key is set), and the original bytes
are deleted. A connection of type `upload` then reads that table like any other source.
"""
import csv
import hashlib
import json
import re
import uuid
import zipfile
from datetime import datetime, timezone
from pathlib import Path

import pyarrow as pa
from sqlalchemy import Column, Integer, MetaData, String, Table, Text, select

from dedupe import crypto
from dedupe.connectors.base import to_string_table
from dedupe.schema import create_and_upgrade

UPLOAD_ID = r"^U[0-9a-f]{12}$"
KINDS = {".csv": "csv", ".tsv": "tsv", ".txt": "csv", ".xlsx": "xlsx"}
MAX_COLUMNS = 300
MAX_ROWS = 2_000_000
MAX_HEADER = 100
MAX_UNZIPPED = 1_000_000_000          # 1 GB of uncompressed Excel parts at most
MAX_RATIO = 200                        # a member that expands more than this is treated as a compression bomb
BATCH = 50_000

_meta = MetaData()
uploads = Table(
    "uploads", _meta,
    Column("id", String(20), primary_key=True),
    Column("filename", String(120), nullable=False),
    Column("kind", String(10), nullable=False),
    Column("rows", Integer, nullable=False),
    Column("columns", Text, nullable=False),
    Column("sha256", String(64), nullable=False),
    Column("uploaded_by", String(200), nullable=False),
    Column("created_at", String(40), nullable=False),
    Column("project_id", String(40)),                  # the project that owns the file: only its connections may use it
)


class UploadError(ValueError):
    pass


def safe_name(name: str) -> str:
    base = re.split(r"[\\/]", name or "")[-1]
    base = re.sub(r"[^A-Za-z0-9 ._()-]", "_", base).strip(" .")
    return base[:120] or "upload"


def _header(cols) -> list[str]:
    names = [("" if c is None else str(c)).strip() for c in cols]
    while names and not names[-1]:
        names.pop()                                     # Excel often reports empty trailing columns
    if not names:
        raise UploadError("the first row must hold the column names")
    if len(names) > MAX_COLUMNS:
        raise UploadError(f"the file has more than {MAX_COLUMNS} columns")
    if any(not n for n in names):
        raise UploadError("every column needs a name in the first row")
    if any(len(n) > MAX_HEADER for n in names):
        raise UploadError(f"a column name is longer than {MAX_HEADER} characters")
    if len(set(names)) != len(names):
        raise UploadError("two columns have the same name")
    return names


def _check_zip(path: Path) -> None:
    try:
        with zipfile.ZipFile(path) as z:
            infos = z.infolist()
            if len(infos) > 10_000:
                raise UploadError("the Excel file has too many parts")
            total = 0
            for i in infos:
                total += i.file_size
                if i.compress_size and i.file_size / max(i.compress_size, 1) > MAX_RATIO:
                    raise UploadError("the Excel file looks like a compression bomb (it expands far more than a spreadsheet would); it was refused")
            if total > MAX_UNZIPPED:
                raise UploadError("the Excel file is too large once unpacked")
            if "xl/workbook.xml" not in {i.filename for i in infos}:
                raise UploadError("this is not an Excel workbook (.xlsx)")
    except zipfile.BadZipFile:
        raise UploadError("this is not a valid Excel workbook (.xlsx)") from None


def _sniff(path: Path, filename: str) -> str:
    with open(path, "rb") as f:
        head = f.read(65536)
    ext = Path(filename).suffix.lower()
    if head.startswith(b"PK\x03\x04"):
        if ext not in (".xlsx", ""):
            raise UploadError("the file content is an Excel/zip file but its name says otherwise")
        return "xlsx"
    if ext == ".xls" or head.startswith(b"\xd0\xcf\x11\xe0"):
        raise UploadError("old .xls files are not supported: save the sheet as .xlsx or CSV")
    if ext not in KINDS or KINDS[ext] == "xlsx":
        raise UploadError("upload a .csv, .tsv or .xlsx file")
    if b"\x00" in head:
        raise UploadError("the file is not text")
    try:
        (head if len(head) < 65536 else head[:65000]).decode("utf-8-sig")       # a multi-byte character may be cut at the very end
    except UnicodeDecodeError:
        raise UploadError("the file is not UTF-8 text: save it as 'CSV UTF-8'") from None
    return KINDS[ext]


def _text_batches(path: Path, kind: str):
    delim = "\t" if kind == "tsv" else ","
    with open(path, encoding="utf-8-sig", newline="") as f:
        reader = csv.reader(f, delimiter=delim)
        try:
            names = _header(next(reader))
        except StopIteration:
            raise UploadError("the file is empty") from None
        yield names
        rows = []
        for r in reader:
            if not any(v.strip() for v in r):
                continue
            rows.append([(v if v != "" else None) for v in (r + [None] * len(names))[:len(names)]])
            if len(rows) >= BATCH:
                yield to_string_table(names, rows)
                rows = []
        if rows:
            yield to_string_table(names, rows)


def _excel_batches(path: Path):
    _check_zip(path)
    from openpyxl import load_workbook
    fh = open(path, "rb")                       # a handle, not a name: the stored file has no .xlsx extension
    wb = load_workbook(fh, read_only=True, data_only=True)
    try:
        it = wb.active.iter_rows(values_only=True)
        try:
            names = _header(next(it))
        except StopIteration:
            raise UploadError("the sheet is empty") from None
        yield names
        rows = []
        for r in it:
            vals = list(r[:len(names)]) + [None] * max(0, len(names) - len(r))
            if all(v is None or str(v).strip() == "" for v in vals):
                continue
            rows.append(vals)
            if len(rows) >= BATCH:
                yield to_string_table(names, rows)
                rows = []
        if rows:
            yield to_string_table(names, rows)
    finally:
        wb.close()
        fh.close()


def _now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


class Uploads:
    def __init__(self, engine, folder: Path):
        self.engine, self.folder = engine, Path(folder)
        create_and_upgrade(engine, _meta)

    def path(self, upload_id: str) -> Path:
        return self.folder / f"{upload_id}.parquet"

    def incoming(self) -> Path:
        """A fresh temporary path for the raw bytes of an upload in progress (inside the workspace, removed after conversion)."""
        d = self.folder / "incoming"
        d.mkdir(parents=True, exist_ok=True)
        return d / uuid.uuid4().hex

    def store(self, raw: Path, filename: str, actor: str, project_id: str | None = None) -> dict:
        raw = Path(raw)
        try:
            name = safe_name(filename)
            kind = _sniff(raw, name)
            digest = hashlib.sha256()
            with open(raw, "rb") as f:
                for chunk in iter(lambda: f.read(1 << 20), b""):
                    digest.update(chunk)
            uid = "U" + uuid.uuid4().hex[:12]
            out = self.path(uid)
            batches = _excel_batches(raw) if kind == "xlsx" else _text_batches(raw, kind)
            try:
                names = next(batches)
            except UploadError:
                raise
            except Exception as e:  # noqa: BLE001
                raise UploadError(f"the file could not be read ({type(e).__name__})") from None
            schema = pa.schema([(n, pa.string()) for n in names])
            rows = 0
            try:
                with crypto.open_writer(out, schema) as w:
                    for t in batches:
                        rows += t.num_rows
                        if rows > MAX_ROWS:
                            raise UploadError(f"the file has more than {MAX_ROWS:,} rows")
                        w.write_table(t.cast(schema))
            except UploadError:
                out.unlink(missing_ok=True)
                raise
            except Exception as e:  # noqa: BLE001 - a parser error can quote cell contents: report only its kind
                out.unlink(missing_ok=True)
                raise UploadError(f"the file could not be read ({type(e).__name__})") from None
            if rows == 0:
                out.unlink(missing_ok=True)
                raise UploadError("the file has a header but no rows")
            with self.engine.begin() as c:
                c.execute(uploads.insert().values(id=uid, filename=name, kind=kind, rows=rows, columns=json.dumps(names), sha256=digest.hexdigest(),
                                                  uploaded_by=actor, created_at=_now(), project_id=project_id))
            return {"upload_id": uid, "filename": name, "kind": kind, "rows": rows, "columns": names}
        finally:
            raw.unlink(missing_ok=True)                    # the original bytes never stay on disk

    def list(self, project_id: str | None = None, limit: int = 100) -> list[dict]:
        q = select(uploads).order_by(uploads.c.created_at.desc()).limit(limit)
        if project_id is not None:
            q = q.where(uploads.c.project_id == project_id)
        with self.engine.connect() as c:
            return [{"upload_id": r.id, "filename": r.filename, "kind": r.kind, "rows": r.rows, "columns": json.loads(r.columns), "uploaded_by": r.uploaded_by,
                     "created_at": r.created_at} for r in c.execute(q)]

    def project_of(self, upload_id: str) -> str | None:
        """The project that owns this upload (None when it is unknown or was uploaded before projects owned files)."""
        with self.engine.connect() as c:
            r = c.execute(select(uploads.c.project_id).where(uploads.c.id == upload_id)).first()
        return r.project_id if r else None


__all__ = ["Uploads", "UploadError", "UPLOAD_ID", "safe_name"]
