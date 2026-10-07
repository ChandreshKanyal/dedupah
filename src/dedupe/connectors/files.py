import csv
import json
import re
from pathlib import Path
from typing import Iterator

import pyarrow as pa
import pyarrow.csv as pacsv
import pyarrow.parquet as pq

from .base import Connector, register, to_string_table
from .guard import ConnectorConfigError, check_path
from .records import batches_of, columns_of, need


class _FileConnector(Connector):
    def __init__(self, path: str, **config):
        super().__init__(path=path, **config)
        self.path = check_path(Path(path))

    def test_connection(self) -> None:
        if not self.path.is_file():
            raise FileNotFoundError(f"{self.path} not found")
        self.columns()


@register
class CsvConnector(_FileConnector):
    """CSV / TSV. Everything is read as text so no value is silently reinterpreted (e.g. leading zeros)."""
    types = ("csv", "tsv")

    @property
    def _delimiter(self) -> str:
        return self.config.get("delimiter") or ("\t" if self.config["type"] == "tsv" else ",")

    def columns(self) -> list[str]:
        with self.path.open(encoding=self.config.get("encoding", "utf-8-sig"), newline="") as f:
            return next(csv.reader(f, delimiter=self._delimiter))

    def read_batches(self, batch_size: int = 50_000) -> Iterator[pa.Table]:
        cols = self.columns()
        reader = pacsv.open_csv(
            self.path,
            read_options=pacsv.ReadOptions(encoding=self.config.get("encoding", "utf-8-sig"), block_size=8 << 20),
            parse_options=pacsv.ParseOptions(delimiter=self._delimiter),
            convert_options=pacsv.ConvertOptions(column_types={c: pa.string() for c in cols}, strings_can_be_null=False),
        )
        for batch in reader:
            yield pa.Table.from_batches([batch])


@register
class XlsxConnector(_FileConnector):
    types = ("xlsx", "excel")

    def _sheet(self, wb):
        name = self.config.get("sheet")
        return wb[name] if name else wb.active

    def columns(self) -> list[str]:
        from openpyxl import load_workbook
        wb = load_workbook(self.path, read_only=True)
        try:
            return [str(h) for h in next(self._sheet(wb).values)]
        finally:
            wb.close()

    def read_batches(self, batch_size: int = 50_000) -> Iterator[pa.Table]:
        from openpyxl import load_workbook
        wb = load_workbook(self.path, read_only=True)
        try:
            it = self._sheet(wb).values
            header = [str(h) for h in next(it)]
            buf: list[list] = []
            for row in it:
                buf.append(list(row))
                if len(buf) >= batch_size:
                    yield to_string_table(header, buf)
                    buf = []
            if buf:
                yield to_string_table(header, buf)
        finally:
            wb.close()


@register
class ParquetConnector(_FileConnector):
    types = ("parquet",)

    def columns(self) -> list[str]:
        return pq.ParquetFile(self.path).schema_arrow.names

    def read_batches(self, batch_size: int = 50_000) -> Iterator[pa.Table]:
        for batch in pq.ParquetFile(self.path).iter_batches(batch_size=batch_size):
            yield _as_text(pa.Table.from_batches([batch]))


def _as_text(t: pa.Table) -> pa.Table:
    """Every column as text; nested or binary columns as their text form."""
    cols = {}
    for name in t.column_names:
        c = t[name]
        try:
            cols[name] = c.cast(pa.string())
        except (pa.ArrowInvalid, pa.ArrowNotImplementedError):
            cols[name] = pa.array([None if v is None else str(v) for v in c.to_pylist()], type=pa.string())
    return pa.table(cols)


@register
class XlsConnector(_FileConnector):
    """Old Excel 97-2003 workbooks (.xls), read with xlrd. Date cells are turned into ISO dates (xls stores them as numbers)."""
    types = ("xls",)

    def _rows(self):
        xlrd = need("xlrd", "files")
        book = xlrd.open_workbook(str(self.path), on_demand=True)
        try:
            sh = book.sheet_by_name(self.config["sheet"]) if self.config.get("sheet") else book.sheet_by_index(0)
            for r in range(sh.nrows):
                row = []
                for c in sh.row(r):
                    if c.ctype == xlrd.XL_CELL_DATE:
                        d = xlrd.xldate_as_datetime(c.value, book.datemode)
                        row.append(d.date().isoformat() if d.time().isoformat() == "00:00:00" else d.isoformat(sep=" "))
                    elif c.ctype == xlrd.XL_CELL_NUMBER and float(c.value).is_integer():
                        row.append(int(c.value))
                    elif c.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        row.append(None)
                    else:
                        row.append(c.value)
                yield row
        finally:
            book.release_resources()

    def columns(self) -> list[str]:
        return [str(h) for h in next(self._rows())]

    def read_batches(self, batch_size: int = 50_000) -> Iterator[pa.Table]:
        it = self._rows()
        header = [str(h) for h in next(it)]
        buf = []
        for row in it:
            buf.append((row + [None] * len(header))[:len(header)])
            if len(buf) >= batch_size:
                yield to_string_table(header, buf)
                buf = []
        yield to_string_table(header, buf)


MAX_JSON_BYTES = 512 * 1024 * 1024


def _dig(doc, path: str | None):
    for part in (path or "").split("."):
        if part:
            doc = doc.get(part) if isinstance(doc, dict) else None
    return doc


@register
class JsonConnector(_FileConnector):
    """A JSON file holding a list of records ([{...}, {...}]), or an object with the list inside (records_path, e.g. "data.customers").
    Also JSON Lines (.jsonl: one record per line), read as a stream. Nested fields become dotted columns (address.city)."""
    types = ("json", "jsonl")

    def _records(self):
        enc = self.config.get("encoding") or "utf-8-sig"
        if self.config["type"] == "jsonl":
            with self.path.open(encoding=enc) as f:
                for line in f:
                    if line.strip():
                        rec = json.loads(line)
                        yield rec if isinstance(rec, dict) else {"value": rec}
            return
        if self.path.stat().st_size > MAX_JSON_BYTES:
            raise ConnectorConfigError("this JSON file is too large to read at once: export it as JSON Lines (one record per line)")
        with self.path.open(encoding=enc) as f:
            doc = json.load(f)
        recs = _dig(doc, self.config.get("records_path"))
        if isinstance(recs, dict):
            recs = [recs]
        if not isinstance(recs, list):
            raise ConnectorConfigError("no list of records was found (set records_path to where the list is, e.g. data.customers)")
        for r in recs:
            yield r if isinstance(r, dict) else {"value": r}

    def columns(self) -> list[str]:
        return columns_of(self._records())

    def read_batches(self, batch_size: int = 50_000) -> Iterator[pa.Table]:
        yield from batches_of(self._records(), batch_size, self.columns())


@register
class XmlConnector(_FileConnector):
    """An XML file with one element per record (record_tag, e.g. "Customer"). Child elements become columns, nested ones dotted
    (Address.City), attributes are prefixed with @. Parsed with defusedxml: no external entities, no entity expansion bombs."""
    types = ("xml",)

    def _records(self):
        ET = need("defusedxml.ElementTree", "files")
        tag = self.config.get("record_tag")
        if not tag or not re.fullmatch(r"[A-Za-z_][\w.-]{0,100}", tag):
            raise ConnectorConfigError("record_tag must be the name of the element that holds one record, e.g. Customer")
        for _, elem in ET.iterparse(str(self.path), events=("end",), forbid_dtd=True):
            if elem.tag.rsplit("}", 1)[-1] == tag:
                yield _element(elem)
                elem.clear()

    def columns(self) -> list[str]:
        return columns_of(self._records())

    def read_batches(self, batch_size: int = 50_000) -> Iterator[pa.Table]:
        yield from batches_of(self._records(), batch_size, self.columns())


def _element(elem) -> dict:
    out: dict = {"@" + k.rsplit("}", 1)[-1]: v for k, v in elem.attrib.items()}
    for child in elem:
        name = child.tag.rsplit("}", 1)[-1]
        value = _element(child) if len(child) or child.attrib else (child.text or "").strip() or None
        if name in out:                                 # a repeated element: keep every value
            prev = out[name]
            out[name] = (prev if isinstance(prev, list) else [prev]) + [value]
        else:
            out[name] = value
    return out


_WIDTH = re.compile(r"^\s*([A-Za-z_][\w ]{0,60}?)\s*:\s*(\d{1,4})\s*$")


@register
class FixedWidthConnector(_FileConnector):
    """Fixed-width text, as mainframes and core banking systems export it. widths lists each column as name:width in order, separated by
    commas ("cust_id:10, name:40, dob:8"); skip_rows skips header lines. Values are trimmed."""
    types = ("fixed_width",)

    def _layout(self) -> list[tuple[str, int, int]]:
        parts = [p for p in str(self.config.get("widths") or "").split(",") if p.strip()]
        if not parts:
            raise ConnectorConfigError("widths must list the columns as name:width, e.g. cust_id:10, name:40")
        out, start = [], 0
        for p in parts:
            m = _WIDTH.match(p)
            if not m:
                raise ConnectorConfigError("each entry of widths must look like name:width (letters, digits, underscore; a width in characters)")
            out.append((m[1].strip(), start, start + int(m[2])))
            start += int(m[2])
        if len({n for n, _, _ in out}) != len(out):
            raise ConnectorConfigError("two columns in widths have the same name")
        return out

    def columns(self) -> list[str]:
        return [n for n, _, _ in self._layout()]

    def read_batches(self, batch_size: int = 50_000) -> Iterator[pa.Table]:
        layout = self._layout()
        skip = int(self.config.get("skip_rows") or 0)
        buf = []
        with self.path.open(encoding=self.config.get("encoding") or "utf-8-sig", newline="") as f:
            for i, line in enumerate(f):
                line = line.rstrip("\r\n")
                if i < skip or not line.strip():
                    continue
                buf.append([line[a:b].strip() or None for _, a, b in layout])
                if len(buf) >= batch_size:
                    yield to_string_table(self.columns(), buf)
                    buf = []
        yield to_string_table(self.columns(), buf)


@register
class AvroConnector(_FileConnector):
    """Apache Avro container files (the schema travels inside the file)."""
    types = ("avro",)

    def _records(self):
        fastavro = need("fastavro", "files")
        with self.path.open("rb") as f:
            yield from fastavro.reader(f)

    def columns(self) -> list[str]:
        return columns_of(self._records())

    def read_batches(self, batch_size: int = 50_000) -> Iterator[pa.Table]:
        yield from batches_of(self._records(), batch_size, self.columns())


@register
class OrcConnector(_FileConnector):
    """Apache ORC files (Hadoop / Hive exports). pyarrow reads ORC on Linux and macOS, not on Windows."""
    types = ("orc",)

    def _file(self):
        return need("pyarrow.orc", "files").ORCFile(str(self.path))

    def columns(self) -> list[str]:
        return self._file().schema.names

    def read_batches(self, batch_size: int = 50_000) -> Iterator[pa.Table]:
        f = self._file()
        for i in range(f.nstripes):
            t = _as_text(pa.Table.from_batches([f.read_stripe(i)]))
            for off in range(0, max(t.num_rows, 1), batch_size):
                yield t.slice(off, batch_size)


FILE_TYPES = {"csv": CsvConnector, "tsv": CsvConnector, "xlsx": XlsxConnector, "xls": XlsConnector, "json": JsonConnector, "jsonl": JsonConnector,
              "xml": XmlConnector, "parquet": ParquetConnector, "avro": AvroConnector, "fixed_width": FixedWidthConnector, "orc": OrcConnector}
FILE_SUFFIX = {"csv": (".csv", ".txt"), "tsv": (".tsv", ".tab", ".txt"), "xlsx": (".xlsx",), "xls": (".xls",), "json": (".json",), "jsonl": (".jsonl", ".ndjson"),
               "xml": (".xml",), "parquet": (".parquet",), "avro": (".avro",), "fixed_width": (".txt", ".dat", ".prn"), "orc": (".orc",)}
FORMAT_OPTIONS = {"delimiter", "encoding", "sheet", "records_path", "record_tag", "widths", "skip_rows"}


def read_local(fmt: str, path: Path, options: dict) -> Connector:
    """A file reader for a file the SERVER chose (a download into the workspace), so it skips the data-folder check."""
    cls = FILE_TYPES.get(fmt)
    if cls is None:
        raise ConnectorConfigError("unknown file format; choose one of: " + ", ".join(sorted(FILE_TYPES)))
    obj = cls.__new__(cls)
    Connector.__init__(obj, type=fmt, path=str(path), **{k: v for k, v in options.items() if k in FORMAT_OPTIONS and v not in (None, "")})
    obj.path = Path(path)
    return obj


MAX_FILES = 5000


@register
class FolderConnector(Connector):
    """Every file matching a pattern in a folder (daily drops, one file per branch): read one after another as a single source. Columns are the
    union of all files' columns; a column missing from a file is empty there. Adds `source_file` with each row's file name."""
    types = ("folder",)

    def __init__(self, **config):
        super().__init__(**config)
        self.folder = check_path(Path(config.get("path", "")))
        pattern = config.get("pattern") or "*"
        if ".." in pattern or pattern.startswith(("/", "\\")) or ":" in pattern:
            raise ConnectorConfigError("pattern must be a file pattern inside the folder, e.g. *.csv or 2026-*/customers.csv")
        self.pattern, self.fmt = pattern, config.get("format") or "csv"
        if self.fmt not in FILE_TYPES:
            raise ConnectorConfigError("unknown file format; choose one of: " + ", ".join(sorted(FILE_TYPES)))

    def _files(self) -> list[Path]:
        if not self.folder.is_dir():
            raise FileNotFoundError("folder not found")
        files = sorted(p for p in self.folder.glob(self.pattern) if p.is_file())
        if len(files) > MAX_FILES:
            raise ConnectorConfigError(f"more than {MAX_FILES} files match the pattern")
        return [check_path(p) for p in files]

    def _readers(self):
        return [(p, read_local(self.fmt, p, self.config)) for p in self._files()]

    def test_connection(self) -> None:
        if not self._files():
            raise FileNotFoundError("no file matches the pattern")
        self.columns()

    def columns(self) -> list[str]:
        cols: dict[str, None] = {}
        for _, r in self._readers():
            for c in r.columns():
                cols.setdefault(c, None)
        return list(cols) + ["source_file"]

    def read_batches(self, batch_size: int = 50_000) -> Iterator[pa.Table]:
        cols = self.columns()
        for p, r in self._readers():
            for t in r.read_batches(batch_size):
                n = t.num_rows
                yield pa.table({c: (t[c] if c in t.column_names else pa.nulls(n, pa.string())) if c != "source_file" else pa.array([p.name] * n, pa.string())
                                for c in cols})
